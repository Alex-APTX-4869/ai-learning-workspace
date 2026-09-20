"""生成非私人解析测试样本；使用文档技能提供的 bundled Python 运行。"""
from pathlib import Path
import argparse
import zipfile

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def build(output: Path):
    output.mkdir(parents=True, exist_ok=False)
    image = Image.new("RGB", (600, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.line([(60, 30), (60, 250), (550, 250)], fill="black", width=3)
    for x, height, label in [(130, 80, "A  8"), (280, 150, "B  15"), (430, 120, "C  12")]:
        draw.rectangle((x, 250-height, x+55, 250), fill="#7895aa")
        draw.text((x, 260), label, fill="black", font_size=18)
    image.save(output / "chart.png")

    doc = Document()
    doc.sections[0].page_width, doc.sections[0].page_height = Inches(8.5), Inches(11)
    for style in ("Title", "Heading 1", "Heading 2", "Normal"):
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[style].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph("Materials parser acceptance", style="Title")
    doc.add_paragraph("This sample checks whether text, table rows, equations and figures remain distinguishable after parsing.")
    doc.add_heading("HTTP requests", level=1)
    doc.add_paragraph("FastAPI 使用路由接收请求。HTTP 422 indicates validation failure.")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Method", "Purpose"
    for left, right in [("GET", "Read a course"), ("POST", "Create a course")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
    doc.add_heading("Fraction and chart", level=1)
    doc.add_paragraph("The native Word fraction below must not become the integer 12.")
    doc.add_paragraph()._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'
    ))
    doc.add_picture(str(output / "chart.png"), width=Inches(4.5))
    doc.save(output / "structure.docx")

    pdf = canvas.Canvas(str(output / "lecture.pdf"), pagesize=(612, 792))
    pdf.setTitle("Materials parser acceptance")
    pdf.bookmarkPage("http")
    pdf.addOutlineEntry("HTTP requests", "http")
    pdf.setFont("Helvetica-Bold", 19)
    pdf.drawString(50, 730, "HTTP requests")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(50, 696, "FastAPI uses routes. HTTP 422 indicates validation failure.")
    pdf.drawString(50, 669, "GET reads a course. POST creates a course.")
    pdf.drawString(55, 615, "1")
    pdf.line(50, 608, 70, 608)
    pdf.drawString(55, 592, "2")
    pdf.drawImage(str(output / "chart.png"), 50, 315, width=450, height=225)
    pdf.showPage()
    pdf.bookmarkPage("database")
    pdf.addOutlineEntry("Database operations", "database")
    pdf.setFont("Helvetica-Bold", 19)
    pdf.drawString(50, 730, "Database operations")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(50, 695, "SQL transactions keep related changes consistent.")
    pdf.save()

    scan = canvas.Canvas(str(output / "scan.pdf"), pagesize=(612, 792))
    scan.drawImage(str(output / "chart.png"), 50, 430, width=500, height=250)
    scan.save()
    writer = PdfWriter()
    writer.append(PdfReader(output / "lecture.pdf"))
    writer.encrypt("fixture-only")
    writer.write(str(output / "encrypted.pdf"))
    (output / "damaged.pdf").write_bytes(b"%PDF-1.7\nnot a document")
    (output / "fake.pdf").write_text("not a PDF")
    # 只有正确 OLE 签名，用于测试明确提示转换器缺失，不冒充有效 DOC 成功样本。
    (output / "legacy-signature.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504))
    with zipfile.ZipFile(output / "structure.docx") as original:
        for name, extra in [("path-traversal.docx", {"../escape.txt": b"no"}),
                            ("macro.docx", {"word/vbaProject.bin": b"no"})]:
            with zipfile.ZipFile(output / name, "w") as modified:
                for info in original.infolist():
                    modified.writestr(info.filename, original.read(info))
                for path, data in extra.items():
                    modified.writestr(path, data)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    build(parser.parse_args().output)
